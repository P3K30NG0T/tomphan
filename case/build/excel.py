"""Builds downloads/pharmacy-deserts-data.xlsx: every dataset behind Acts 1-4, with formulas for derived numbers.
Run after case/build/build.py:  python3 case/build/excel.py  (then recalculate with LibreOffice to cache values).
Blue = input or source data, black = formula. Each sheet names its source in row 2."""
import csv, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

d = json.load(open("data/case.json"))
facts = json.load(open("case/build/facts.json"))
S = facts["sources"]
a1, a2, a3, a4 = d["act1"], d["act2"], d["act3"], d["act4"]


def F(**k): return Font(**{"name": "Arial", "size": 10, **k})
BLUE, BOLD, H1, NOTE = F(color="0000FF"), F(bold=True), F(bold=True, size=12), F(color="666666", italic=True)
HEAD, YEL = PatternFill("solid", fgColor="DDE3EA"), PatternFill("solid", fgColor="FFF2CC")
INT, USD, USDM, PCT, PCT1, DEC = '#,##0;(#,##0);-', '$#,##0;($#,##0);-', '$#,##0.0;($#,##0.0);-', '0%;(0%);-', '0.0%;(0.0%);-', '#,##0.00'
wb = Workbook()
first = True


def sheet(name, title, source):
    global first
    ws = wb.active if first else wb.create_sheet()
    first = False
    ws.title = name
    ws["A1"], ws["A2"] = title, source
    ws["A1"].font, ws["A2"].font = H1, NOTE
    return ws


def header(ws, row, cols, widths=None):
    for j, h in enumerate(cols, start=1):
        c = ws.cell(row=row, column=j, value=h); c.font = BOLD; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def put(ws, r, c, v, fmt=None, font=None):
    cell = ws.cell(row=r, column=c, value=v)
    if fmt: cell.number_format = fmt
    if font: cell.font = font
    elif not (isinstance(v, str) and v.startswith("=")) and isinstance(v, (int, float)): cell.font = BLUE
    return cell


# ---------- README
ws = sheet("README", "Pharmacy Deserts · Weis Markets: data behind Acts 1-4", f"Tom Phan (tomphann.vercel.app). Built {d['built']}. Independent analysis on public data; no inside information; not investment advice.")
rows = [
    ("How to read", "Blue numbers are source data or inputs you can change. Black numbers are formulas. Row 2 of every sheet names the source."),
    ("A1_US_Quarterly", "Act 1. US pharmacies and drug stores by quarter (BLS QCEW), with change since Q3 2023."),
    ("A1_States_Quarterly", "Act 1. The same count for Weis's seven states."),
    ("A1_State_Change", "Act 1. Change by state from Q3 2023 to the latest quarter."),
    ("A1_Demand", "Act 1. Pharmacies vs Medicare Part D 30-day prescriptions, 2014 onward, indexed."),
    ("A1_Timeline", "Act 1. Dated events with sources."),
    ("A2_Counties", "Act 2. All 365 counties in the seven states: pharmacies, people, age, cars, health, Weis pharmacies."),
    ("A2_State_Demand", "Act 2. Medicare prescriptions per stand-alone pharmacy by state, 2024."),
    ("A3_Payroll", "Act 3. Pay and jobs per pharmacy, 2014-2025 (BLS QCEW)."),
    ("A3_Calculator", "Act 3. One stand-alone pharmacy's year: change the blue inputs."),
    ("A3_Benchmarks", "Act 3. Industry benchmarks from NCPA, Walgreens, CVS and BLS wage data."),
    ("A4_Weis_Financials", "Act 4. Weis Markets financials from SEC filings."),
    ("A4_Weis_Pharmacies", "Act 4. Pharmacies registered to Weis in the NPI registry, with county."),
    ("Sources", "Every source with its link."),
    ("Note on counts", "BLS counts businesses whose main activity is selling medicine (chain drugstores and independents). Pharmacy counters inside supermarkets and big-box stores are counted under the store's main business, so they are not included. County totals are slightly below state totals because BLS withholds some county detail."),
]
header(ws, 4, ["Sheet", "What it holds"], [24, 120])
for i, (a, b) in enumerate(rows, start=5):
    ws.cell(row=i, column=1, value=a).font = BOLD
    ws.cell(row=i, column=2, value=b).alignment = Alignment(wrap_text=True)

# ---------- Act 1
ws = sheet("A1_US_Quarterly", "US pharmacies and drug stores, by quarter", f"Source: {S['qcew']['title']}, private sector. {S['qcew']['url']}")
header(ws, 4, ["Quarter", "Pharmacies", "Jobs", "Change in pharmacies vs Q3 2023", "% change vs Q3 2023"], [12, 14, 14, 22, 18])
us = a1["us_quarterly"]
base_row = 5 + [r["q"] for r in us].index(a1["base"])
for i, r in enumerate(us, start=5):
    put(ws, i, 1, r["q"]); put(ws, i, 2, r["estabs"], INT); put(ws, i, 3, r["jobs"], INT)
    put(ws, i, 4, f"=B{i}-$B${base_row}", INT); put(ws, i, 5, f"=B{i}/$B${base_row}-1", PCT1)
ws.cell(row=base_row, column=1).fill = YEL

ws = sheet("A1_States_Quarterly", "Pharmacies and drug stores in Weis's seven states, by quarter", f"Source: {S['qcew']['title']}. {S['qcew']['url']}")
sts = list(a1["states_quarterly"].keys())
header(ws, 4, ["Quarter"] + sts + ["Seven states"], [12] + [10] * len(sts) + [14])
qs = [r["q"] for r in a1["states_quarterly"][sts[0]]]
for i, q in enumerate(qs, start=5):
    put(ws, i, 1, q)
    for j, s in enumerate(sts, start=2):
        put(ws, i, j, next(r["estabs"] for r in a1["states_quarterly"][s] if r["q"] == q), INT)
    put(ws, i, len(sts) + 2, f"=SUM(B{i}:{get_column_letter(len(sts) + 1)}{i})", INT)

ws = sheet("A1_State_Change", f"Change in pharmacies, {a1['base']} to {a1['last']}", f"Source: {S['qcew']['title']}; Weis store counts from the fiscal 2025 10-K.")
header(ws, 4, ["State", "Weis stores", f"Pharmacies {a1['base']}", f"Pharmacies {a1['last']}", "Change", "% change", f"Jobs {a1['base']}", f"Jobs {a1['last']}", "% change in jobs"], [16, 11, 14, 14, 10, 10, 12, 12, 14])
rows = sorted(a1["states_change"], key=lambda r: r["latest"] / r["base"])
for i, r in enumerate(rows, start=5):
    put(ws, i, 1, r["name"]); put(ws, i, 2, r["weis"], INT); put(ws, i, 3, r["base"], INT); put(ws, i, 4, r["latest"], INT)
    put(ws, i, 5, f"=D{i}-C{i}", INT); put(ws, i, 6, f"=D{i}/C{i}-1", PCT1)
    put(ws, i, 7, r["jobs_base"], INT); put(ws, i, 8, r["jobs_latest"], INT); put(ws, i, 9, f"=H{i}/G{i}-1", PCT1)
n = 4 + len(rows)
t = n + 1
put(ws, t, 1, "Seven states", font=BOLD)
for c in (2, 3, 4, 7, 8):
    put(ws, t, c, f"=SUM({get_column_letter(c)}5:{get_column_letter(c)}{n})", INT)
put(ws, t, 5, f"=D{t}-C{t}", INT); put(ws, t, 6, f"=D{t}/C{t}-1", PCT1); put(ws, t, 9, f"=H{t}/G{t}-1", PCT1)
u = t + 1
put(ws, u, 1, "United States", font=BOLD); put(ws, u, 3, a1["us_change"]["base"], INT); put(ws, u, 4, a1["us_change"]["latest"], INT)
put(ws, u, 5, f"=D{u}-C{u}", INT); put(ws, u, 6, f"=D{u}/C{u}-1", PCT1)
put(ws, u, 7, a1["us_change"]["jobs_base"], INT); put(ws, u, 8, a1["us_change"]["jobs_latest"], INT); put(ws, u, 9, f"=H{u}/G{u}-1", PCT1)
put(ws, u + 2, 1, "Rite Aid stores at its first filing (Oct 2023)"); put(ws, u + 2, 4, a1["context"]["rite_aid_stores_at_filing"], INT)
put(ws, u + 3, 1, "Share of the US drop that cannot be Rite Aid (at least)"); put(ws, u + 3, 4, f"=1-D{u + 2}/-E{u}", PCT)

ws = sheet("A1_Demand", "Pharmacies vs Medicare prescriptions", f"Sources: {S['qcew']['title']} (annual averages); {S['partd']['title']} (national 30-day fills, summed over all drugs). {S['partd']['url']}")
header(ws, 4, ["Year", "Pharmacies (annual avg)", "Medicare Part D 30-day fills", "Medicare Part D claims", "Pharmacies index (2014=100)", "Fills index (2014=100)", "Medicare fills per pharmacy"], [8, 18, 22, 20, 18, 16, 18])
pd = {r["year"]: r for r in a1["partd"]}
for i, r in enumerate(a1["us_annual"], start=5):
    put(ws, i, 1, r["year"]); put(ws, i, 2, r["estabs"], INT)
    if r["year"] in pd:
        put(ws, i, 3, pd[r["year"]]["fills"], INT); put(ws, i, 4, pd[r["year"]]["claims"], INT)
        put(ws, i, 6, f"=C{i}/$C$5*100", "0.0"); put(ws, i, 7, f"=C{i}/B{i}", INT)
    put(ws, i, 5, f"=B{i}/$B$5*100", "0.0")

ws = sheet("A1_Timeline", "What happened, with sources", "Sources: see the link in each row.")
header(ws, 4, ["When", "What happened", "Why it matters", "Source"], [11, 90, 40, 60])
for i, r in enumerate(a1["timeline"], start=5):
    ws.cell(row=i, column=1, value=r["label"]); ws.cell(row=i, column=2, value=r["event"]).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=i, column=3, value=r["why"]).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=i, column=4, value=S[r["src"]]["url"])

# ---------- Act 2
ws = sheet("A2_Counties", "All counties in Weis's seven states", f"Sources: {S['qcew']['title']} (pharmacies, {a1['base']} and {a1['last']}); {S['svi']['title']} (people, age, cars, poverty, insurance); {S['places']['title']} (diabetes, high blood pressure, lack of transport; 2023 estimates); {S['npi']['title']}.")
cols = ["FIPS", "County", "State", "Residents", "Residents 65+", "% 65+", "% households with no car", "% below 150% of poverty", "% uninsured",
        f"Pharmacies {a1['base']}", f"Pharmacies {a1['last']}", "Change", "% change", "Pharmacies per 10,000 residents", "Residents per pharmacy", "Residents 65+ per pharmacy",
        "% adults with diabetes", "% adults with high blood pressure", "% lacking reliable transport", "Weis pharmacies (NPI)"]
header(ws, 4, cols, [8, 18, 6, 12, 12, 7, 10, 10, 9, 11, 11, 8, 8, 12, 12, 12, 10, 12, 11, 10])
for i, c in enumerate(sorted(a2["counties"], key=lambda x: (x["state"], x["name"])), start=5):
    vals = [c["fips"], c["name"], c["state"], c["pop"], c["age65"], (c["pct65"] or 0) / 100, (c["noveh"] or 0) / 100, (c["pov150"] or 0) / 100, (c["uninsured"] or 0) / 100, c["ph_base"], c["ph_latest"]]
    fmts = [None, None, None, INT, INT, PCT1, PCT1, PCT1, PCT1, INT, INT]
    for j, (v, f) in enumerate(zip(vals, fmts), start=1):
        put(ws, i, j, v, f, font=None if j > 3 else F())
    put(ws, i, 12, f"=K{i}-J{i}", INT); put(ws, i, 13, f'=IF(J{i}=0,"",K{i}/J{i}-1)', PCT)
    put(ws, i, 14, f"=K{i}/D{i}*10000", DEC); put(ws, i, 15, f'=IF(K{i}=0,"none",D{i}/K{i})', INT); put(ws, i, 16, f'=IF(K{i}=0,"none",E{i}/K{i})', INT)
    for j, k in ((17, "diabetes"), (18, "bphigh"), (19, "lacktrpt")):
        put(ws, i, j, (c[k] / 100) if c[k] is not None else None, PCT1)
    put(ws, i, 20, c["weis_rx"], INT)
ws.auto_filter.ref = f"A4:T{4 + len(a2['counties'])}"

ws = sheet("A2_State_Demand", "Medicare prescriptions per stand-alone pharmacy, 2024", f"Sources: {S['partd']['title']} (state level, 2024); {S['qcew']['title']} (2024 average of quarterly counts); {S['svi']['title']}.")
header(ws, 4, ["State", "Weis stores", "Medicare 30-day fills 2024", "Pharmacies 2024 (avg)", "Fills per pharmacy", "Residents", "Residents 65+", "Residents 65+ per pharmacy"], [16, 11, 20, 16, 16, 14, 14, 18])
for i, r in enumerate(a2["state_demand"], start=5):
    put(ws, i, 1, r["name"]); put(ws, i, 2, r["weis"], INT); put(ws, i, 3, r["fills_2024"], INT); put(ws, i, 4, r["ph_2024"], INT)
    put(ws, i, 5, f"=C{i}/D{i}", INT); put(ws, i, 6, r["pop"], INT); put(ws, i, 7, r["age65"], INT); put(ws, i, 8, f"=G{i}/D{i}", INT)
k = 5 + len(a2["state_demand"])
put(ws, k, 1, "United States", font=BOLD); put(ws, k, 5, a2["us_fills_per_ph_2024"], INT)

# ---------- Act 3
ws = sheet("A3_Payroll", "Pay and jobs per pharmacy, United States", f"Source: {S['qcew']['title']}, annual averages.")
header(ws, 4, ["Year", "Pharmacies", "Jobs", "Total pay ($)", "Average pay per job ($)", "Pay per pharmacy ($)", "Jobs per pharmacy", "Pay per pharmacy vs 2014"], [8, 12, 12, 18, 16, 16, 12, 16])
for i, r in enumerate(a1["us_annual"], start=5):
    put(ws, i, 1, r["year"]); put(ws, i, 2, r["estabs"], INT); put(ws, i, 3, r["jobs"], INT); put(ws, i, 4, r["wages"], USD); put(ws, i, 5, r["avg_pay"], USD)
    put(ws, i, 6, f"=D{i}/B{i}", USD); put(ws, i, 7, f"=C{i}/B{i}", "0.0"); put(ws, i, 8, f"=F{i}/$F$5-1", PCT1)

ws = sheet("A3_Calculator", "One stand-alone pharmacy's year (change the blue inputs)", "Defaults: see column C. Rent and other costs are an assumption, not a published figure.")
cd = a3["calc_defaults"]
header(ws, 4, ["Item", "Value", "Where the default comes from"], [42, 16, 100])
inp = [("Prescriptions per day", cd["rx_per_day"], INT, cd["rx_per_day_src"]),
       ("Revenue per prescription ($)", cd["revenue_per_rx"], '$#,##0.00', cd["revenue_per_rx_src"]),
       ("Gross margin (share of revenue kept after the drug cost)", cd["gross_margin"], PCT1, cd["gross_margin_src"]),
       ("Payroll per year ($)", cd["payroll"], USD, cd["payroll_src"]),
       ("Rent and other costs per year ($)", cd["other_costs"], USD, cd["other_costs_src"]),
       ("Days open per year", 365, INT, "Assumption: open every day")]
for i, (a, v, f, s) in enumerate(inp, start=5):
    ws.cell(row=i, column=1, value=a); put(ws, i, 2, v, f); ws.cell(row=i, column=3, value=s)
out = [("Gross profit per prescription ($)", "=B6*B7", '$#,##0.00'),
       ("Prescriptions per year", "=B5*B10", INT),
       ("Revenue per year ($)", "=B13*B6", USD),
       ("Gross profit per year ($)", "=B13*B12", USD),
       ("Fixed costs per year ($)", "=B8+B9", USD),
       ("Profit before tax ($)", "=B15-B16", USD),
       ("Profit as share of revenue", "=B17/B14", PCT1),
       ("Prescriptions per day needed to break even", "=B16/B12/B10", "0"),
       ("Profit change if paid $1 less per prescription ($)", "=-B13", USD)]
for i, (a, f, fm) in enumerate(out, start=12):
    ws.cell(row=i, column=1, value=a).font = BOLD if "Profit" in a and "share" not in a else F(); put(ws, i, 2, f, fm)
ws.cell(row=22, column=1, value="Profit before tax ($) by prescriptions per day (rows) and revenue per prescription (columns)").font = BOLD
rx_list, rev_list = [100, 130, 150, 180, 210, 250], [70, 74, 78.6, 82, 86]
for j, v in enumerate(rev_list, start=2):
    put(ws, 23, j, v, '$#,##0.00')
for i, rx in enumerate(rx_list, start=24):
    put(ws, i, 1, rx, INT)
    for j in range(2, 2 + len(rev_list)):
        col = get_column_letter(j)
        put(ws, i, j, f"=$A{i}*$B$10*{col}$23*$B$7-$B$16", USD)

ws = sheet("A3_Benchmarks", "Industry benchmarks", "Sources: NCPA Digest (2024, 2025, 2026 editions), Walgreens fiscal 2024 results and 10-K, CVS fiscal 2025 10-K and Q4 2025 results, BLS OEWS May 2025.")
header(ws, 4, ["Measure", "Value", "Year", "Source"], [62, 16, 10, 70])
rows = []
for n in a3["ncpa"]:
    rows += [(f"Independent pharmacies (count)", n["stores"], INT, n["year"], S[n["src"]]["url"]),
             (f"Prescriptions per independent pharmacy per year", n["rx_per_store"], INT, n["year"], S[n["src"]]["url"]),
             (f"Independent pharmacy market ($ billions)", n["market_bn"], "0.0", n["year"], S[n["src"]]["url"])]
    if n["gross_margin"]:
        rows.append(("Independent pharmacy gross margin", n["gross_margin"], PCT1, n["year"], S[n["src"]]["url"]))
for w in a3["walgreens"]:
    rows += [(f"Walgreens US Retail Pharmacy sales ($ billions)", w["sales_bn"], "0.0", f"FY{w['fy']}", S["wba1200"]["url"]),
             (f"Walgreens US Retail Pharmacy adjusted operating income ($ billions)", w["adj_op_income_bn"], "0.00", f"FY{w['fy']}", S["wba1200"]["url"])]
w24 = a3["walgreens"][1]
rows += [("Walgreens prescriptions, 30-day adjusted (billions)", w24["rx_30day_bn"], "0.0", "FY2024", S["wba10k"]["url"]),
         ("Walgreens US locations", w24["locations"], INT, "FY2024", S["wba10k"]["url"]),
         ("CVS prescriptions, 30-day equivalent (billions)", a3["cvs_2025"]["rx_30day_bn"], "0.0", "2025", S["cvs10k"]["url"]),
         ("CVS share of US retail prescriptions", a3["cvs_2025"]["retail_share"], PCT1, "2025", S["cvs10k"]["url"]),
         ("CVS stores closed in 2025", a3["cvs_2025"]["closed_2025"], INT, "2025", S["cvs10k"]["url"])]
for occ in ("pharmacist", "technician"):
    for area, v in a3["oews"][occ].items():
        rows.append((f"Mean annual wage, {occ}, {area}", v, USD, "2025", S["oews"]["url"]))
for i, (a, v, f, y, s) in enumerate(rows, start=5):
    ws.cell(row=i, column=1, value=a); put(ws, i, 2, v, f); ws.cell(row=i, column=3, value=str(y)); ws.cell(row=i, column=4, value=s)
r = 5 + len(rows) + 1
ws.cell(row=r, column=1, value="Walgreens adjusted operating margin, FY2023").font = BOLD; put(ws, r, 2, "=B{}/B{}".format(5 + [x[0] for x in rows].index("Walgreens US Retail Pharmacy adjusted operating income ($ billions)"), 5 + [x[0] for x in rows].index("Walgreens US Retail Pharmacy sales ($ billions)")), PCT1)
ws.cell(row=r + 1, column=1, value="Walgreens adjusted operating margin, FY2024").font = BOLD
i_s = 5 + len(rows) - 1 - [x[0] for x in rows][::-1].index("Walgreens US Retail Pharmacy sales ($ billions)")
i_o = 5 + len(rows) - 1 - [x[0] for x in rows][::-1].index("Walgreens US Retail Pharmacy adjusted operating income ($ billions)")
put(ws, r + 1, 2, f"=B{i_o}/B{i_s}", PCT1)
ws.cell(row=r + 2, column=1, value="Walgreens adjusted operating income per 30-day prescription, FY2024 ($)").font = BOLD
i_rx = 5 + [x[0] for x in rows].index("Walgreens prescriptions, 30-day adjusted (billions)")
put(ws, r + 2, 2, f"=B{i_o}/B{i_rx}", '$#,##0.00')

# ---------- Act 4
ws = sheet("A4_Weis_Financials", "Weis Markets: selected financials ($ millions)", f"Sources: {S['weis10k']['title']}; {S['weisxbrl']['title']}; {S['weisq2']['title']}. {a4['fin_note']}")
years = [r["year"] for r in a4["fin"]]
header(ws, 4, ["Item"] + [str(y) for y in years], [44] + [12] * len(years))
items = [("Net sales", "net_sales", USDM), ("Pharmacy share of net sales", "pharmacy_share", PCT1), ("Operating income", "op_income", USDM),
         ("Net income", "net_income", USDM), ("Gross margin", "gross_margin", PCT1), ("Cash and cash equivalents", "cash", USDM),
         ("Marketable securities", "securities", USDM), ("Cash from operations", "ocf", USDM), ("Capital spending", "capex", USDM), ("Depreciation and amortization", "da", USDM)]
rowof = {}
for i, (lab, key, f) in enumerate(items, start=5):
    ws.cell(row=i, column=1, value=lab); rowof[key] = i
    for j, r in enumerate(a4["fin"], start=2):
        if r.get(key) is not None:
            put(ws, i, j, r[key], f)
calc = [("Pharmacy sales (derived)", lambda c: f'=IF({c}{rowof["net_sales"]}="","",{c}{rowof["net_sales"]}*{c}{rowof["pharmacy_share"]})', USDM),
        ("Operating margin", lambda c: f'=IF({c}{rowof["net_sales"]}="","",{c}{rowof["op_income"]}/{c}{rowof["net_sales"]})', PCT1),
        ("Cash plus marketable securities", lambda c: f"={c}{rowof['cash']}+{c}{rowof['securities']}", USDM),
        ("Cash from operations minus capital spending", lambda c: f"={c}{rowof['ocf']}-{c}{rowof['capex']}", USDM)]
for i, (lab, fn, f) in enumerate(calc, start=6 + len(items)):
    ws.cell(row=i, column=1, value=lab).font = BOLD
    for j in range(2, 2 + len(years)):
        put(ws, i, j, fn(get_column_letter(j)), f)
h = a4["h1_2026"]
r0 = 8 + len(items) + len(calc)
ws.cell(row=r0, column=1, value="First half of 2026").font = H1
for i, (lab, v, f) in enumerate([("Revenue ($ billions)", h["revenue_bn"], "0.00"), ("Revenue growth vs first half 2025", h["revenue_growth"], PCT1),
                                 ("Operating income", h["op_income"], USDM), ("Net income", h["net_income"], USDM),
                                 ("Pharmacy revenue lost to Medicare negotiated prices", h["ira_pharmacy_revenue_hit"], USDM),
                                 ("Cash and cash equivalents (June 27, 2026)", h["cash"], USDM), ("Marketable securities (June 27, 2026)", h["securities"], USDM)], start=r0 + 1):
    ws.cell(row=i, column=1, value=lab); put(ws, i, 2, v, f)

ws = sheet("A4_Weis_Pharmacies", "Pharmacies registered to Weis Markets", f"Source: {S['npi']['title']} ({S['npi']['url']}); county from {S['geocoder']['title']}. Active registrations; a registration can outlive a closed counter.")
header(ws, 4, ["NPI", "City", "State", "County FIPS", "County"], [14, 22, 7, 12, 30])
cn = {c["fips"]: c["name"] for c in a2["counties"]}
for i, r in enumerate(sorted(a4["weis_pharmacies"], key=lambda x: (x["state"], x["city"])), start=5):
    ws.cell(row=i, column=1, value=r["npi"]); ws.cell(row=i, column=2, value=r["city"]); ws.cell(row=i, column=3, value=r["state"])
    ws.cell(row=i, column=4, value=r["county_fips"]); ws.cell(row=i, column=5, value=cn.get(r["county_fips"], ""))
k = 5 + len(a4["weis_pharmacies"])
ws.cell(row=k + 1, column=1, value="Total").font = BOLD; put(ws, k + 1, 3, f"=COUNTA(A5:A{k - 1})", INT)

# ---------- Sources
ws = sheet("Sources", "Sources", "All links checked on " + d["built"])
header(ws, 4, ["Key", "Source", "Link"], [12, 80, 100])
for i, (k, v) in enumerate(S.items(), start=5):
    ws.cell(row=i, column=1, value=k); ws.cell(row=i, column=2, value=v["title"]); ws.cell(row=i, column=3, value=v["url"])

for w in wb.worksheets:
    w.sheet_view.showGridLines = False
wb.save("downloads/pharmacy-deserts-data.xlsx")
print("saved")
