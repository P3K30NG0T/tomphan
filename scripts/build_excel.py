"""Build downloads/tomphan-h100-model.xlsx from data/market.json (the research harness output).
Run from the repo root: python3 scripts/build_excel.py  (then recalculate with LibreOffice to cache values).
Blue = input, black = formula, green = link to another sheet. Every hardcoded number carries its source."""
import json
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

mk = json.load(open("data/market.json", encoding="utf-8"))
A = {k: v["value"] for k, v in mk["assumptions"].items()}

def F(**k): return Font(**{"name": "Arial", "size": 10, **k})
BLUE, BOLD, GREEN, H1 = F(color="0000FF"), F(bold=True), F(color="008000"), F(bold=True, size=12)
YEL, HEAD = PatternFill("solid", fgColor="FFFF00"), PatternFill("solid", fgColor="DDE3EA")
USD, USD2, PCT = '$#,##0;($#,##0);-', '$#,##0.00;($#,##0.00);-', '0.0%;(0.0%);-'

wb = Workbook()

# ---------------- Ladder (evidence for the decline assumption)
lad = wb.active; lad.title = "Ladder"
lad["A1"] = "Price ladder: rent vs years on the market (evidence for the rent-decline assumption)"; lad["A1"].font = H1
lad["A2"] = f"Source: GetDeploying GPU price dataset (CC BY 4.0), on-demand median, average of the last 4 complete weeks to {mk['asOf']}. Availability year = analyst input."
lad["A4"] = "As-of year (decimal)"; asof = date.fromisoformat(mk["asOf"]); lad["B4"] = round(asof.year + (asof.month - 1) / 12, 2); lad["B4"].font = BLUE
for j, h in enumerate(["GPU", "Volume availability (year)", "Years on market", "Rent $/GPU-hr", "LN(rent)", "12-month change"]):
    c = lad.cell(row=6, column=1 + j, value=h); c.font = BOLD; c.fill = HEAD
avail = {"A100": 2020.4, "H100": 2022.9, "H200": 2024.3, "B200": 2025.1, "B300": 2025.8}
rows = [m for m in mk["models"] if m["label"] in avail]
for i, m in enumerate(rows, start=7):
    lad.cell(row=i, column=1, value=m["label"])
    c = lad.cell(row=i, column=2, value=avail[m["label"]]); c.font = BLUE
    lad.cell(row=i, column=3, value=f"=$B$4-B{i}").number_format = "0.0"
    c = lad.cell(row=i, column=4, value=m["latest"]); c.font = BLUE; c.number_format = USD2
    lad.cell(row=i, column=5, value=f"=LN(D{i})").number_format = "0.000"
    c = lad.cell(row=i, column=6, value=m["change"]); c.font = BLUE; c.number_format = PCT
last = 6 + len(rows)
for r, (l, f, fmt) in enumerate([("Slope of LN(rent) per year", f"=SLOPE(E7:E{last},C7:C{last})", "0.0000"),
                                 ("R²", f"=RSQ(E7:E{last},C7:C{last})", "0.00"),
                                 ("Rent lost per year of age", f"=1-EXP(B{last + 2})", PCT)], start=last + 2):
    lad.cell(row=r, column=1, value=l).font = BOLD
    c = lad.cell(row=r, column=2, value=f); c.number_format = fmt; c.font = BOLD
LAD_DECLINE = f"Ladder!$B${last + 4}"
lad.column_dimensions["A"].width = 30
for col in "BCDEF": lad.column_dimensions[col].width = 16

# ---------------- Inputs
ws = wb.create_sheet("Inputs", 0)
ws["A1"] = "tomphan. Should you buy an H100 today to rent it out? (one GPU)"; ws["A1"].font = H1
ws["A2"] = "Blue = input. Green = linked from another sheet. Black = formula. Yellow = scenario selector. USD per GPU unless noted."
ws["A4"] = "Active scenario"; ws["A4"].font = BOLD
ws["B4"] = "Base"; ws["B4"].font = BLUE; ws["B4"].fill = YEL
dv = DataValidation(type="list", formula1='"Bull,Base,Bear"'); ws.add_data_validation(dv); dv.add("B4")
ws["C4"] = '=MATCH(B4,{"Bull","Base","Bear"},0)'
for j, h in enumerate(["Assumption", "Active", "Bull", "Base", "Bear", "Unit", "Source / why"]):
    c = ws.cell(row=6, column=1 + j, value=h); c.font = BOLD; c.fill = HEAD
why = {k: v["why"] for k, v in mk["assumptions"].items()}
spec = [  # key, label, base, bull, bear, unit, fmt, source
    ("capex", "Capex per GPU, all-in", A["capex"], "=D7", "=D7", "USD", USD, why["capex"]),
    ("price", "Rent, year 1", A["price"], "=D8", "=D8", "USD/GPU-hr", USD2, "SemiAnalysis H100 1-yr rental index, Mar 2026. " + why["price"]),
    ("lock", "Years locked by contract", A["lock"], 3, "=D9", "years", "0", why["lock"]),
    ("decline", "Rent decline per year", f"=MIN(0.4,MAX(0.08,ROUND({LAD_DECLINE},2)))", "=MAX(0.05,D10-0.08)", "=MIN(0.45,D10+0.08)", "% per yr", PCT, "Linked from the Ladder sheet: fitted rate rounded to 2 decimals, bounded 8–40%, the same rule the harness applies."),
    ("util", "Utilization", A["util"], "=MIN(0.95,D11+0.05)", "=MAX(0.4,D11-0.15)", "% of hours", PCT, why["util"]),
    ("kw", "Power per GPU incl. cooling", 1.4, "=D12", "=D12", "kW", "0.00", "Analyst input: 700W GPU + server share, PUE ~1.3."),
    ("elec", "Electricity", A["elec"], "=D13", "=D13", "USD/kWh", "$0.000", "EIA Monthly Energy Review Table 9.8, US industrial, May 2026."),
    ("otherOpex", "Other opex", 0.30, "=D14", "=D14", "USD per available GPU-hr", USD2, "Analyst input: colocation, network, staff, maintenance."),
    ("life", "Economic life", A["life"], 6, 4, "years", "0", why["life"]),
    ("residual", "Resale value", 0.05, "=D16", "=D16", "% of capex", PCT, "Analyst input."),
    ("wacc", "Cost of capital", A["wacc"], "=D17", "=D17", "%", PCT, why["wacc"]),
    ("tax", "Tax rate", 0.21, "=D18", "=D18", "%", PCT, "US federal corporate rate; losses assumed usable by the wider business."),
    ("hrs", "Hours per year", 8760, "=D19", "=D19", "hours", "#,##0", "")
]
cell = {}
for r, (k, label, base, bull, bear, unit, fmt, src) in enumerate(spec, start=7):
    ws.cell(row=r, column=1, value=label)
    ws.cell(row=r, column=2, value=f"=INDEX(C{r}:E{r},$C$4)").number_format = fmt
    for j, v in enumerate((bull, base, bear)):
        c = ws.cell(row=r, column=3 + j, value=v); c.number_format = fmt
        c.font = GREEN if isinstance(v, str) and "Ladder!" in v else (F() if isinstance(v, str) else BLUE)
    ws.cell(row=r, column=6, value=unit); ws.cell(row=r, column=7, value=src)
    cell[k] = f"Inputs!$B${r}"
ws.column_dimensions["A"].width = 28; ws.column_dimensions["G"].width = 95
for col in "BCDEF": ws.column_dimensions[col].width = 13

# ---------------- Model
m = wb.create_sheet("Model", 1)
m["A1"] = "Cash flows per GPU (USD)"; m["A1"].font = H1
m["A2"] = "Rent holds at the year-1 rate while locked by contract, then follows the market curve: rent × (1 − decline)^(year − 1). Columns past the life show zero."
labels = ["Year", "Active (1 = in life)", "Rent ($/hr)", "Revenue", "Power", "Other opex", "EBITDA", "Depreciation", "EBIT", "Tax (negative = shield)", "Resale value", "Free cash flow", "Cumulative FCF", "Rent coefficient (PV per $1 of rent)", "Capex coefficient (PV per $1 of capex)"]
for i, l in enumerate(labels): m.cell(row=4 + i, column=1, value=l).font = BOLD if i in (0, 6, 11) else F()
cols = [chr(ord("B") + i) for i in range(9)]
c_ = cell
for t, col in enumerate(cols):
    m[f"{col}4"] = str(t); m[f"{col}4"].font = BOLD; m[f"{col}4"].fill = HEAD
    if t == 0:
        m[f"{col}15"] = f"=-{c_['capex']}"; m[f"{col}16"] = f"={col}15"; m[f"{col}18"] = -1; continue
    m[f"{col}5"] = f"=IF({t}<={c_['life']},1,0)"
    m[f"{col}6"] = f"=IF({t}<={c_['lock']},{c_['price']},{c_['price']}*(1-{c_['decline']})^({t}-1))*{col}5"
    m[f"{col}7"] = f"={col}6*{c_['hrs']}*{c_['util']}"
    m[f"{col}8"] = f"=-{c_['kw']}*{c_['hrs']}*{c_['util']}*{c_['elec']}*{col}5"
    m[f"{col}9"] = f"=-{c_['otherOpex']}*{c_['hrs']}*{col}5"
    m[f"{col}10"] = f"={col}7+{col}8+{col}9"
    m[f"{col}11"] = f"=-{c_['capex']}/{c_['life']}*{col}5"
    m[f"{col}12"] = f"={col}10+{col}11"
    m[f"{col}13"] = f"=-{col}12*{c_['tax']}"
    m[f"{col}14"] = f"=IF({t}={c_['life']},{c_['residual']}*{c_['capex']},0)"
    m[f"{col}15"] = f"={col}10+{col}13+{col}14"
    m[f"{col}16"] = f"={cols[t - 1]}16+{col}15"
    m[f"{col}17"] = f"=IF({c_['price']}>0,{col}6/{c_['price']},0)*{c_['hrs']}*{c_['util']}*(1-{c_['tax']})/(1+{c_['wacc']})^{t}"
    m[f"{col}18"] = f"=({col}5*{c_['tax']}/{c_['life']}+IF({t}={c_['life']},{c_['residual']},0))/(1+{c_['wacc']})^{t}"
for r in range(5, 19):
    for col in cols:
        c = m[f"{col}{r}"]
        if c.value is not None: c.number_format = "0" if r == 5 else (USD2 if r == 6 else ("0.000" if r >= 17 else USD))
m.column_dimensions["A"].width = 36
for col in cols: m.column_dimensions[col].width = 12
m["A21"] = "Results"; m["A21"].font = H1
res = [("NPV per GPU", f"=B15+NPV({c_['wacc']},C15:J15)", USD),
       ("NPV, 1,000-GPU cluster", "=B22*1000", USD),
       ("IRR", "=IFERROR(IRR(B15:J15),\"n/a\")", PCT),
       ("Payback (years)", '=IFERROR(MATCH(TRUE,INDEX(C16:J16>=0,0),0)-1+(-INDEX(B16:I16,MATCH(TRUE,INDEX(C16:J16>=0,0),0))/INDEX(C15:J15,MATCH(TRUE,INDEX(C16:J16>=0,0),0))),"Not within life")', "0.00"),
       ("Break-even year-1 rent ($/hr)", f"={c_['price']}-B22/SUM(C17:J17)", USD2),
       ("Max capex per GPU to break even", f"={c_['capex']}-B22/SUM(B18:J18)", USD),
       ("Cash cost per billed GPU-hour", f"={c_['kw']}*{c_['elec']}+{c_['otherOpex']}/{c_['util']}", USD2),
       ("Year-1 EBIT margin, economic life", "=IFERROR(C12/C7,0)", PCT),
       ("Year-1 EBIT margin, 6-year accounting life", f"=IFERROR((C10-{c_['capex']}/6)/C7,0)", PCT)]
for i, (l, f, fmt) in enumerate(res, start=22):
    m.cell(row=i, column=1, value=l); c = m.cell(row=i, column=2, value=f); c.number_format = fmt; c.font = BOLD
m["C26"] = "NPV is linear in rent and in capex, so the break-evens are exact: current value minus NPV ÷ PV coefficient."

# ---------------- H100 weekly data
d = wb.create_sheet("H100 weekly")
d["A1"] = "H100 on-demand median rent, USD/GPU-hr (GetDeploying, CC BY 4.0)"; d["A1"].font = H1
d["A3"] = "Week"; d["B3"] = "Median rent"; d["A3"].font = d["B3"].font = BOLD
h = next(x for x in mk["models"] if x["label"] == "H100")
for i, (wk, v) in enumerate(h["series"], start=4):
    d.cell(row=i, column=1, value=wk); c = d.cell(row=i, column=2, value=v); c.font = BLUE; c.number_format = USD2
d.column_dimensions["A"].width = 14; d.column_dimensions["B"].width = 14

# ---------------- Sources
s = wb.create_sheet("Sources")
s["A1"] = "Sources"; s["A1"].font = H1
src = [("GetDeploying GPU price dataset (CC BY 4.0)", "https://getdeploying.com/dataset/gpu-prices"),
       ("SemiAnalysis H100 1-yr rental index", mk["manual"]["contract1y"]["url"]),
       ("EIA Monthly Energy Review, electricity prices", mk["manual"]["power"]["url"]),
       ("CoreWeave Form 10-K FY2025", mk["manual"]["coreweave"]["url"]),
       ("Research harness code and run log", "https://github.com/P3K30NG0T/tomphan")]
for i, (t, u) in enumerate(src, start=3): s.cell(row=i, column=1, value=t); s.cell(row=i, column=2, value=u)
s["A10"] = f"Data as of {mk['asOf']}. Independent analysis on open data by Phan Nguyen Hong Quang (tomphan.). Not investment advice."
s.column_dimensions["A"].width = 48; s.column_dimensions["B"].width = 100

for w in wb.worksheets:
    for row in w.iter_rows():
        for c in row:
            if c.value is not None and (c.font is None or c.font.name != "Arial"):
                c.font = Font(name="Arial", size=10, bold=c.font.b, color=c.font.color)
wb.save("downloads/tomphan-h100-model.xlsx")
print("saved")
